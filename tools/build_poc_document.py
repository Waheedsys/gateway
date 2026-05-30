import argparse
import os
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DEFAULT_OUT = Path("docs/AI_Gateway_Elasticsearch_MCP_POC.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        set_cell_shading(hdr[i], "F2F4F7")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_footer(doc):
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("AI Gateway POC")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def parse_args():
    parser = argparse.ArgumentParser(description="Build the AI Gateway POC DOCX.")
    parser.add_argument("--out", default=os.getenv("DOCGEN_OUTPUT", str(DEFAULT_OUT)))
    parser.add_argument("--unique", action="store_true", help="Append a timestamp to the output filename.")
    return parser.parse_args()


def unique_path(path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return path.with_name(f"{path.stem}_{stamp}{path.suffix}")


def main():
    args = parse_args()
    out = Path(args.out)
    if args.unique:
        out = unique_path(out)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    add_footer(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("AI Gateway POC")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    r = subtitle.add_run("Elasticsearch Context Engineering and MCP Tool Calling")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    meta = doc.add_paragraph()
    meta.add_run("Purpose: ").bold = True
    meta.add_run("Demonstrate a Go-based LLM gateway that stores memory in Elasticsearch, retrieves context during inference, calls tools through MCP, and exposes observability through Kibana.")

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "This proof of concept replaces PostgreSQL persistence with Elasticsearch-backed document storage and retrieval. "
        "The gateway stores conversations, messages, and inference logs as Elasticsearch documents. During inference, it builds a model prompt from recent conversation history, relevant retrieved context, and optional MCP tool results."
    )
    doc.add_paragraph(
        "The current POC validates the core platform pattern: Elasticsearch acts as conversational memory and retrieval infrastructure, while MCP provides standardized access to external tools such as file reading, directory listing, or future business-system integrations."
    )

    doc.add_heading("2. Architecture", level=1)
    add_table(
        doc,
        ["Component", "Role", "Current Implementation"],
        [
            ["Go Gateway", "Receives API requests, builds context, calls the LLM provider, and streams or returns responses.", "cmd/gateway with Chi routes"],
            ["Elasticsearch", "Stores conversations, messages, inference logs, and supports context retrieval.", "Indexes: conversations, messages, inference_logs"],
            ["Redis", "Supports rate limiting and short-lived response caching.", "Existing Redis service"],
            ["MCP Server", "Exposes external tools through a standard protocol.", "Built-in test server under cmd/testmcp"],
            ["LLM Provider", "Generates assistant responses from assembled prompt context.", "OpenRouter or Anthropic provider abstraction"],
            ["Kibana", "Visual UI for inspecting Elasticsearch documents and logs.", "Added as docker-compose service on port 5601"],
        ],
        [1.35, 3.05, 2.1],
    )

    doc.add_heading("3. POC Demo Flow", level=1)
    for step in [
        "Start Elasticsearch, Redis, Kibana, and the Go gateway.",
        "Create a conversation using POST /conversations.",
        "Send a memory-setting inference request, for example: My preferred search engine is Elasticsearch. Remember this.",
        "Ask a follow-up question such as: What search engine do I prefer?",
        "Verify that the assistant answers from stored context.",
        "Create a file under C:/tmp and call the MCP read_file tool through /infer.",
        "Open Kibana at http://localhost:5601 and inspect messages, inference_logs, and raw_metadata.mcp.",
    ]:
        add_number(doc, step)

    doc.add_heading("4. Working Capabilities", level=1)
    for item in [
        "Conversation creation and retrieval through the existing REST API.",
        "Message persistence in Elasticsearch.",
        "Inference log persistence in Elasticsearch with provider metadata and token usage.",
        "Context assembly from recent messages and Elasticsearch full-text retrieval.",
        "MCP tool listing through GET /mcp/tools.",
        "Direct MCP tool calls through POST /mcp/tools/call.",
        "Inline MCP calls during inference using tool:<name> JSON arguments.",
        "Kibana visibility for conversations, messages, inference logs, and MCP metadata.",
        "SSE streaming remains available for inference responses.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("5. Validation Evidence", level=1)
    doc.add_paragraph(
        "The POC has already demonstrated successful context recall. After storing the statement that Elasticsearch is the preferred database, a later inference request asked which database was preferred, and the assistant answered that it was Elasticsearch."
    )
    doc.add_paragraph(
        "The MCP path has also been validated. A request using tool:read_file with C:/tmp/mcp-test.txt returned an assistant response containing the file content: Hello from MCP server. This confirms that the gateway detected the inline tool call, invoked the MCP server, injected the result into the prompt, and saved the assistant response."
    )

    doc.add_heading("6. Feature Development Roadmap", level=1)
    add_table(
        doc,
        ["Priority", "Feature", "Why It Matters"],
        [
            ["P0", "AUTH_ENABLED and RATE_LIMIT_ENABLED toggles", "Allows local testing without removing production security code."],
            ["P0", "Persistent MCP sessions", "Avoids starting a new MCP process per request and improves latency/reliability."],
            ["P1", "Automatic tool planner", "Allows users to ask naturally while the gateway decides whether to call MCP tools."],
            ["P1", "Context deduplication", "Prevents the current message or repeated retrieved text from appearing twice in prompts."],
            ["P1", "Vector search with embeddings", "Improves semantic retrieval beyond keyword matching."],
            ["P2", "Tool allow/deny policy", "Restricts risky tools and supports safer production deployments."],
            ["P2", "Kibana dashboards", "Adds ready-made views for latency, tokens, errors, and MCP usage."],
            ["P2", "Provider-native tool schemas", "Bridges MCP tools into provider-native tool/function calling where supported."],
        ],
        [0.75, 2.0, 3.75],
    )

    doc.add_heading("7. Production Considerations", level=1)
    for item in [
        "Enable Elasticsearch security, API keys, snapshots, and production index lifecycle policies.",
        "Run Elasticsearch with replicas in a multi-node environment. Yellow health is expected only in local single-node development.",
        "Store secrets outside .env for deployed environments.",
        "Add request tracing and correlation IDs across gateway, MCP tool calls, and provider responses.",
        "Add tests around repository behavior, MCP tool errors, prompt construction, and inference logging.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("8. Recommended Next Sprint", level=1)
    for step in [
        "Add environment toggles for authentication and rate limiting.",
        "Convert MCP from per-request process startup to persistent managed sessions.",
        "Add an LLM-based tool planner for automatic MCP selection.",
        "Deduplicate retrieved context and add prompt-size controls.",
        "Add Elasticsearch vector embeddings for semantic context retrieval.",
        "Create a Kibana dashboard export for demo and operations.",
    ]:
        add_number(doc, step)

    try:
        doc.save(out)
    except PermissionError:
        out = unique_path(out)
        doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
