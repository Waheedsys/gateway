import argparse
import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DEFAULT_OUTPUT_DIR = Path("docs/generated")


def slugify(value):
    value = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")
    return value[:70] or "document"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, width in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(width))
        node.set(qn("w:type"), "dxa")


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_footer(doc, title):
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(title[:80])
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    if subtitle:
        sub = doc.add_paragraph()
        sub.paragraph_format.space_after = Pt(12)
        r = sub.add_run(subtitle)
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_metadata_table(doc, values):
    rows = [(k, v) for k, v in values.items() if v]
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row_idx, (key, value) in enumerate(rows):
        cells = table.rows[row_idx].cells
        cells[0].text = key
        cells[1].text = str(value)
        cells[0].width = Inches(1.45)
        cells[1].width = Inches(5.05)
        set_cell_shading(cells[0], "F2F4F7")
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
    doc.add_paragraph()


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(str(text))


def normalize_sections(data):
    sections = data.get("sections") or []
    normalized = []
    if isinstance(sections, str):
        sections = [s.strip() for s in sections.split(",") if s.strip()]
    for section in sections:
        if isinstance(section, str):
            normalized.append({"heading": section, "body": ""})
        elif isinstance(section, dict):
            normalized.append(
                {
                    "heading": section.get("heading") or section.get("title") or "Section",
                    "body": section.get("body") or section.get("content") or "",
                    "bullets": section.get("bullets") or [],
                }
            )
    if normalized:
        return normalized
    return [
        {"heading": "Overview", "body": ""},
        {"heading": "Key Capabilities", "body": ""},
        {"heading": "Implementation Approach", "body": ""},
        {"heading": "Risks and Considerations", "body": ""},
        {"heading": "Next Steps", "body": ""},
    ]


def default_body(topic, heading, doc_type, audience):
    templates = {
        "Overview": f"This section introduces {topic} and frames why it matters for {audience or 'the intended audience'}. The focus is on practical value, implementation fit, and the decisions needed to move forward.",
        "Key Capabilities": f"{topic} should be evaluated by the concrete capabilities it enables, the operational workflows it improves, and the measurable outcomes it can support.",
        "Implementation Approach": f"A practical implementation should start with a narrow working path, validate the core workflow, then expand into production concerns such as reliability, security, observability, and maintainability.",
        "Risks and Considerations": f"Important considerations include integration complexity, operational ownership, access control, data quality, monitoring, and clear rollback paths.",
        "Next Steps": f"The recommended next step is to agree on success criteria, choose a small pilot scope, assign owners, and validate the workflow with realistic inputs.",
    }
    return templates.get(
        heading,
        f"This section covers {heading.lower()} for {topic}. It is written as part of a {doc_type or 'business'} document for {audience or 'the target audience'}.",
    )


def build_document(data, out_dir):
    topic = str(data.get("topic") or data.get("title") or "Generated Document").strip()
    doc_type = str(data.get("doc_type") or data.get("type") or "brief").strip()
    audience = str(data.get("audience") or "").strip()
    tone = str(data.get("tone") or "professional").strip()
    title = str(data.get("title") or topic).strip()
    subtitle = str(data.get("subtitle") or f"{doc_type.title()} for {audience}" if audience else doc_type.title())

    doc = Document()
    configure_styles(doc)
    add_footer(doc, title)
    add_title(doc, title, subtitle)
    add_metadata_table(
        doc,
        {
            "Topic": topic,
            "Document Type": doc_type,
            "Audience": audience,
            "Tone": tone,
            "Generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        },
    )

    summary = data.get("summary") or data.get("executive_summary")
    if not summary:
        summary = (
            f"This {doc_type} summarizes {topic} for {audience or 'stakeholders'}. "
            "It provides a structured overview, key considerations, and actionable next steps in a concise professional format."
        )
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(str(summary))

    for section in normalize_sections(data):
        heading = section["heading"]
        doc.add_heading(heading, level=1)
        body = section.get("body") or default_body(topic, heading, doc_type, audience)
        doc.add_paragraph(str(body))
        for bullet in section.get("bullets") or []:
            add_bullet(doc, bullet)

    recommendations = data.get("recommendations") or data.get("next_steps") or []
    if isinstance(recommendations, str):
        recommendations = [s.strip() for s in recommendations.split(";") if s.strip()]
    if recommendations:
        doc.add_heading("Recommended Actions", level=1)
        for item in recommendations:
            p = doc.add_paragraph(style="List Number")
            p.add_run(str(item))

    out_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{slugify(title)}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    out = out_dir / filename
    doc.save(out)
    return out


def parse_args():
    parser = argparse.ArgumentParser(description="Generate a general DOCX from JSON input.")
    parser.add_argument("--json", help="JSON payload string.")
    parser.add_argument("--input", help="Path to JSON payload file.")
    parser.add_argument("--out-dir", default=str(DEFAULT_OUTPUT_DIR))
    return parser.parse_args()


def main():
    args = parse_args()
    if args.input:
        data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    elif args.json:
        data = json.loads(args.json)
    else:
        data = {"topic": "Generated Document"}
    out = build_document(data, Path(args.out_dir))
    print(out)


if __name__ == "__main__":
    main()
